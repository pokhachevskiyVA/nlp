# -*- coding: utf-8 -*-
"""Плеядный анализ по кластерам + регрессионные отчёты с формулами.
Опирается на spearman_core.load_and_clean (та же очистка данных)."""
import re
import numpy as np
import pandas as pd
from scipy import stats
import matplotlib
import matplotlib.pyplot as plt
from itertools import combinations
from openpyxl import Workbook
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
from openpyxl.utils import get_column_letter

import spearman_core as sc
from spearman_core import load_and_clean, POINT_PREFIX, effect_word
from glossary import describe
from report_text import write_methodology

matplotlib.rcParams["font.family"] = "DejaVu Sans"

# ---------- цели (результат) ----------
def point_columns(df):
    """Столбцы-результаты (универсально: точки/дистанции/ручной выбор)."""
    cols = sc.result_columns(df)
    if cols:
        return cols
    return [c for c in df.columns if str(c).lower().startswith(POINT_PREFIX.lower())]


HIGH_TARGET_MARKERS = ("PC1", "Композит", "Сумма")


def target_is_high(target_name):
    """Цель уже ориентирована 'больше=лучше'? (PC1/композит/сумма — да)."""
    return any(k in str(target_name) for k in HIGH_TARGET_MARKERS)


def effect_for(r, target_name):
    """Словесный эффект показателя с учётом смысла цели и ориентации результата."""
    if target_is_high(target_name):
        return "улучшает" if r > 0 else "ухудшает"
    return effect_word(r)          # сырой результат (точка/дистанция) — по ориентации


def build_targets(df):
    """dict {название: Series}: PC1 (общий результат) + сумма/композит.
    Обе цели ориентированы БОЛЬШЕ = ЛУЧШЕ (учитывая orientation: время → знак развёрнут)."""
    pts = point_columns(df)
    orient = sc.orientation()
    P = df[pts].apply(lambda s: s.fillna(s.median()))
    X = P.values.astype(float)
    Xs = (X - X.mean(0)) / (X.std(0, ddof=0) + 1e-12)
    if orient == "low":                 # время: меньше=лучше → инвертируем в «успех»
        Xs = -Xs
    U, S, Vt = np.linalg.svd(Xs, full_matrices=False)
    pc1 = U[:, 0] * S[0]
    comp = Xs.mean(axis=1)              # композит-успех (выше=лучше)
    if np.corrcoef(pc1, comp)[0, 1] < 0:
        pc1 = -pc1
    expl = (S ** 2 / (S ** 2).sum())[0]
    pc1 = pd.Series(pc1, index=df.index, name="PC1 (общий результат)")
    if orient == "low":
        second = pd.Series(comp, index=df.index, name="Композит результата (выше=лучше)")
    else:
        second = df[pts].sum(axis=1, min_count=1).rename("Сумма результатов")
    return {pc1.name: pc1, second.name: second}, expl


def make_clusters(target):
    """Медианное деление игроков: 'сильные' (>= медианы) / 'слабые' (< медианы)."""
    med = target.median()
    lab = pd.Series(np.where(target >= med, "сильные", "слабые"), index=target.index)
    return lab


# ---------- отбор показателей ----------
def candidate_indicators(df, target, mask=None, min_cov=0.6, leak_thr=0.95):
    """Числовые показатели (не точки, не цель) с достаточным покрытием данными.
    Исключаются 'производные' столбцы, почти идеально совпадающие с целью
    (|r|>=leak_thr) — это функции от точек (напр. заранее посчитанная сумма),
    иначе анализ выродится в тавтологию."""
    pts = set(point_columns(df))
    idx = df.index if mask is None else df.index[mask]
    t = target.loc[idx]
    summ = df.loc[idx, list(pts)].sum(axis=1, min_count=1)  # суммарные попадания
    cols = []
    for c in df.columns:
        if c in pts:
            continue
        s = df.loc[idx, c]
        if s.notna().sum() < max(3, int(min_cov * len(idx))) or s.nunique(dropna=True) <= 1:
            continue
        # утечка: столбец почти совпадает с целью ИЛИ с суммой попаданий
        # (агрегаты бросков вроде заранее посчитанной суммы — не предикторы)
        leak = False
        for ref in (t, summ):
            m = s.notna() & ref.notna()
            if m.sum() >= 4 and s[m].nunique() > 1 and ref[m].nunique() > 1:
                r = stats.spearmanr(s[m], ref[m])[0]
                if not np.isnan(r) and abs(r) >= leak_thr:
                    leak = True
                    break
        if leak:
            continue
        cols.append(c)
    return cols


def rank_by_target(df, target, cols, mask=None, method="spearman"):
    """|корреляция| каждого показателя с целью (в пределах подвыборки)."""
    idx = df.index if mask is None else df.index[mask]
    t = target.loc[idx]
    rows = []
    for c in cols:
        s = df.loc[idx, c]
        m = s.notna() & t.notna()
        if m.sum() < 4 or s[m].nunique() < 2 or t[m].nunique() < 2:
            continue
        if method == "pearson":
            r, p = stats.pearsonr(s[m], t[m])
        else:
            r, p = stats.spearmanr(s[m], t[m])
        rows.append((c, r, p, int(m.sum())))
    res = pd.DataFrame(rows, columns=["показатель", "r", "p", "n"])
    return res.reindex(res["r"].abs().sort_values(ascending=False).index).reset_index(drop=True)


# ---------- значимость / порог по выборке ----------
def crit_r(n, alpha=0.05):
    """Критическое |r| для значимости при данном объёме выборки n (двусторонний).
    Связь считается существенной, если |r| >= crit_r. Выведено из t-приближения:
    r = t / sqrt(n-2+t^2)."""
    if n is None or n < 4:
        return np.nan
    tc = stats.t.ppf(1 - alpha / 2, n - 2)
    return float(tc / np.sqrt(n - 2 + tc ** 2))


def crit_r_words(n, alpha=0.05):
    rc = crit_r(n, alpha)
    warn = "" if (n is not None and n >= 15) else \
        f"  ⚠ спортсменов мало (n={n}<15) — связи ненадёжны."
    if np.isnan(rc):
        return f"n={n}: слишком мало данных для оценки значимости.{warn}"
    return (f"При n={n} существенными (p<{alpha}) считаются связи с |r| ≥ {rc:.3f}."
            f"{warn}")


# ---------- корреляционные плеяды ----------
def _short(name, n=22):
    s = re.sub(r"\s+", " ", str(name)).strip()
    return s if len(s) <= n else s[: n - 1] + "…"


def _wrap(name, width=15, max_lines=3):
    """Перенос длинного названия узла на несколько строк для читаемости."""
    s = re.sub(r"\s+", " ", str(name).replace("\n", " ")).strip()
    s = re.sub(r"__\d+$", "", s)
    tokens = re.split(r"[ ,]+", s)
    lines, cur = [], ""
    for w in tokens:
        if len(cur) + len(w) + 1 <= width:
            cur = (cur + " " + w).strip()
        else:
            if cur:
                lines.append(cur)
            cur = w
    if cur:
        lines.append(cur)
    if len(lines) > max_lines:
        lines = lines[:max_lines]
        lines[-1] += "…"
    return "\n".join(lines)


def _tshort(name):
    """Короткая подпись цели в центре узла."""
    if "PC1" in str(name):
        return "PC1"
    if "Композит" in str(name) or "Сумма" in str(name):
        return "ИТОГ"
    return _short(name, 10)


def auto_top_k(n):
    """Оптимальное число узлов для читаемости в зависимости от объёма выборки."""
    if n is None:
        return 8
    if n < 22:
        return 6
    if n < 40:
        return 8
    return 9


def pleiad(df, target, cluster_name="все", mask=None, method="spearman",
           threshold=0.4, top_k=None, ax=None, title=None,
           alpha=None, annotate=True):
    """Строит корреляционную плеяду: цель в центре (жёлтая), вокруг —
    top_k показателей, сильнее всего связанных с целью. Рёбра —
    красные (+) / синие (−), толщина ∝ |r|, на ребре подписан коэффициент r.

    Порог связи:
      • alpha не None  → показываются ВСЕ существенные связи (p<alpha); порог |r|
        вычисляется автоматически по выборке (crit_r).
      • alpha is None  → фиксированный порог |r|>=threshold.
    Возвращает (fig, edges_df, node_df)."""
    idx = df.index if mask is None else df.index[mask]
    n_obs = len(idx)
    if top_k is None:
        top_k = auto_top_k(n_obs)
    rc = crit_r(n_obs, alpha) if alpha is not None else threshold
    eff_thr = rc if alpha is not None else threshold

    cols = candidate_indicators(df, target, mask)
    ranked = rank_by_target(df, target, cols, mask, method)
    chosen = ranked.head(top_k)["показатель"].tolist()
    nodes = [target.name] + chosen
    data = pd.concat([target.loc[idx].rename(target.name), df.loc[idx, chosen]], axis=1)

    def corr(a, b):
        m = a.notna() & b.notna()
        if m.sum() < 4 or a[m].nunique() < 2 or b[m].nunique() < 2:
            return np.nan, np.nan, int(m.sum())
        if method == "pearson":
            r, p = stats.pearsonr(a[m], b[m])
        else:
            r, p = stats.spearmanr(a[m], b[m])
        return r, p, int(m.sum())

    def keep(r, p):
        if np.isnan(r):
            return False
        return (p < alpha) if alpha is not None else (abs(r) >= threshold)

    k = len(chosen)
    pos = {target.name: (0.0, 0.0)}
    for i, c in enumerate(chosen):
        ang = 2 * np.pi * i / max(k, 1) + np.pi / 2
        pos[c] = (np.cos(ang), np.sin(ang))

    edges = []
    for a, b in combinations(nodes, 2):
        r, p, n = corr(data[a], data[b])
        if keep(r, p):
            edges.append((a, b, r, p, n))

    if ax is None:
        fig, ax = plt.subplots(figsize=(13, 13))
    else:
        fig = ax.figure

    def _lblpos(a, b):
        # подпись r: у рёбер к цели — на 62% от центра к периферии (не на центре)
        if a == target.name or b == target.name:
            per = b if a == target.name else a
            return pos[per][0] * 0.62, pos[per][1] * 0.62
        mx, my = (pos[a][0] + pos[b][0]) / 2, (pos[a][1] + pos[b][1]) / 2
        # рёбра периферия-периферия, проходящие через центр, уводим наружу
        d = (mx ** 2 + my ** 2) ** 0.5
        if d < 0.5:
            if d < 0.08:                       # почти через центр — сместить перпендикулярно
                dx, dy = pos[b][0] - pos[a][0], pos[b][1] - pos[a][1]
                nrm = (dx ** 2 + dy ** 2) ** 0.5 or 1
                return -dy / nrm * 0.5, dx / nrm * 0.5
            f = 0.55 / d
            return mx * f, my * f
        return mx, my

    for a, b, r, p, n in edges:
        x = [pos[a][0], pos[b][0]]; y = [pos[a][1], pos[b][1]]
        ax.plot(x, y, color=("#c0392b" if r > 0 else "#2471a3"),
                lw=0.7 + 6.5 * abs(r), alpha=0.65, zorder=1, solid_capstyle="round")
        if annotate:
            mx, my = _lblpos(a, b)
            ax.text(mx, my, f"{r:+.2f}", ha="center", va="center",
                    fontsize=8, fontweight="bold", zorder=4,
                    color=("#7b241c" if r > 0 else "#1a5276"),
                    bbox=dict(boxstyle="round,pad=0.12", fc="white",
                              ec="none", alpha=0.85))
    # узлы: маркеры небольшие, подписи вынесены наружу круга
    for node, (x, y) in pos.items():
        is_t = node == target.name
        ax.scatter([x], [y], s=(2800 if is_t else 750),
                   c=("#f1c40f" if is_t else "#aed6f1"),
                   edgecolors="#333", linewidths=1.3, zorder=2)
    ax.text(0, 0, _tshort(target.name), ha="center", va="center",
            fontsize=11, fontweight="bold", zorder=6,
            bbox=dict(boxstyle="circle,pad=0.3", fc="#f1c40f", ec="#333", lw=1.3))
    for node in chosen:
        x, y = pos[node]
        lx, ly = x * 1.34, y * 1.34
        ha = "left" if x > 0.08 else ("right" if x < -0.08 else "center")
        va = "bottom" if y > 0.08 else ("top" if y < -0.08 else "center")
        ax.text(lx, ly, _wrap(node), ha=ha, va=va, fontsize=8.5, zorder=3,
                linespacing=0.95)
    ax.set_xlim(-1.95, 1.95); ax.set_ylim(-1.8, 1.8)
    ax.axis("off")
    if title is None:
        crit_txt = (f"p<{alpha}, |r|≥{rc:.3f}" if alpha is not None
                    else f"|r|≥{threshold}")
        title = (f"Плеяда — {target.name} | кластер: {cluster_name}\n"
                 f"{method}, {crit_txt}, n={n_obs}"
                 + ("" if n_obs >= 15 else "  ⚠ n<15"))
    ax.set_title(title, fontsize=11, fontweight="bold")

    # таблица связей с целью: что улучшает/ухудшает результат
    node_rows = []
    for _, b, r, p, n in [e for e in edges if e[0] == target.name] + \
                          [(e[1], e[0], e[2], e[3], e[4]) for e in edges if e[1] == target.name]:
        cat, desc = describe(b)
        node_rows.append((b, r, effect_for(r, target.name), cat, desc, p, n))
    node_df = pd.DataFrame(node_rows,
                           columns=["показатель", "r с целью", "эффект", "категория",
                                    "что это (для врача)", "p", "n"]) \
        .drop_duplicates("показатель").sort_values("r с целью", key=lambda s: s.abs(), ascending=False)
    edges_df = pd.DataFrame(edges, columns=["узел A", "узел B", "r", "p", "n"])
    return fig, edges_df, node_df.reset_index(drop=True)


# ---------- регрессия с уравнением ----------
def _ols(y, X):
    """OLS через numpy. X без столбца единиц. Возвращает dict с коэф., R², adjR², se, p."""
    n, k = X.shape
    Xd = np.column_stack([np.ones(n), X])
    beta, *_ = np.linalg.lstsq(Xd, y, rcond=None)
    yhat = Xd @ beta
    resid = y - yhat
    ss_res = float(resid @ resid)
    ss_tot = float(((y - y.mean()) ** 2).sum())
    r2 = 1 - ss_res / ss_tot if ss_tot > 0 else np.nan
    dof = n - k - 1
    adj = 1 - (1 - r2) * (n - 1) / dof if dof > 0 else np.nan
    # стандартные ошибки коэффициентов
    if dof > 0:
        sigma2 = ss_res / dof
        try:
            cov = sigma2 * np.linalg.inv(Xd.T @ Xd)
            se = np.sqrt(np.diag(cov))
            with np.errstate(divide="ignore", invalid="ignore"):
                tvals = np.where(se > 0, beta / se, np.nan)
            pvals = 2 * stats.t.sf(np.abs(tvals), dof)
        except np.linalg.LinAlgError:
            se = np.full(k + 1, np.nan); pvals = np.full(k + 1, np.nan)
    else:
        se = np.full(k + 1, np.nan); pvals = np.full(k + 1, np.nan)
    return dict(beta=beta, se=se, p=pvals, r2=r2, adj=adj, n=n, k=k)


def best_regression(df, target, mask=None, method="spearman",
                    n_candidates=8, max_terms=3, min_cov=0.6):
    """Отбирает кандидатов по |r| с целью, перебирает подмножества до max_terms,
    выбирает модель с максимальным adj R². Возвращает результат best + ранжирование."""
    cols = candidate_indicators(df, target, mask, min_cov)
    ranked = rank_by_target(df, target, cols, mask, method)
    cand = ranked.head(n_candidates)["показатель"].tolist()
    idx = df.index if mask is None else df.index[mask]
    best = None
    for size in range(1, min(max_terms, len(cand)) + 1):
        for combo in combinations(cand, size):
            sub = pd.concat([target.loc[idx].rename("__y__"), df.loc[idx, list(combo)]], axis=1).dropna()
            if len(sub) < len(combo) + 3:
                continue
            y = sub["__y__"].values.astype(float)
            X = sub[list(combo)].values.astype(float)
            m = _ols(y, X)
            score = m["adj"]
            if best is None or (not np.isnan(score) and score > best["score"]):
                best = dict(score=score, terms=list(combo), model=m)
    return best, ranked


def equation_str(target_name, terms, model, digits=4):
    b = model["beta"]
    parts = [f"{b[0]:.{digits}g}"]
    for name, coef in zip(terms, b[1:]):
        sign = "+" if coef >= 0 else "−"
        parts.append(f" {sign} {abs(coef):.{digits}g}·[{_short(name,28)}]")
    return f"{target_name} = " + "".join(parts)


# ---------- xlsx-отчёт ----------
HDR = PatternFill("solid", fgColor="FF305496")
HDRF = Font(color="FFFFFFFF", bold=True)
SIG = PatternFill("solid", fgColor="FFE06666")
SIG5 = PatternFill("solid", fgColor="FFF4CCCC")
THIN = Side(style="thin", color="FFD9D9D9")
BORD = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def _hrow(ws, r, vals, start=1):
    for j, v in enumerate(vals, start=start):
        c = ws.cell(r, j, v); c.font = HDRF; c.fill = HDR
        c.alignment = Alignment(horizontal="center", wrap_text=True)


def regression_xlsx(out_path, df, targets, method="spearman"):
    """Отчёт: для каждой цели — общая модель и модели по кластерам.
    Уравнение, R²/adjR², таблица коэффициентов, ранжирование показателей."""
    wb = Workbook(); wb.remove(wb.active)
    used = []
    for tname, target in targets.items():
        clusters = make_clusters(target)
        scopes = [("все игроки", None)]
        for cl in ["сильные", "слабые"]:
            scopes.append((f"кластер: {cl}", (clusters == cl).values))
        ws = wb.create_sheet(_short(tname, 28))
        r = 1
        ws.cell(r, 1, f"Регрессионные модели результата «{tname}» ({method})").font = Font(bold=True, size=12); r += 2
        for scope_name, mask in scopes:
            best, ranked = best_regression(df, target, mask, method)
            if best:
                used += list(best["terms"])
            used += list(ranked.head(10)["показатель"])
            n = int((mask.sum()) if mask is not None else len(df))
            ws.cell(r, 1, f"▶ {scope_name} (n={n})").font = Font(bold=True, size=11); r += 1
            if best is None:
                ws.cell(r, 1, "недостаточно данных для модели"); r += 2; continue
            m = best["model"]
            ws.cell(r, 1, "Уравнение:").font = Font(bold=True)
            ws.cell(r, 2, equation_str(tname, best["terms"], m)); r += 1
            ws.cell(r, 1, "R²:").font = Font(bold=True); ws.cell(r, 2, round(m["r2"], 3))
            ws.cell(r, 3, "adj R²:").font = Font(bold=True); ws.cell(r, 4, round(m["adj"], 3)); r += 1
            # коэффициенты
            _hrow(ws, r, ["Член", "Коэффициент", "Ст.ошибка", "p-value"]); r += 1
            names = ["(константа)"] + best["terms"]
            for i, nm in enumerate(names):
                ws.cell(r, 1, _short(nm, 40)); ws.cell(r, 2, round(float(m["beta"][i]), 5))
                ws.cell(r, 3, round(float(m["se"][i]), 5) if not np.isnan(m["se"][i]) else None)
                pc = ws.cell(r, 4, round(float(m["p"][i]), 4) if not np.isnan(m["p"][i]) else None)
                if not np.isnan(m["p"][i]):
                    if m["p"][i] < 0.01: pc.fill = SIG
                    elif m["p"][i] < 0.05: pc.fill = SIG5
                for j in range(1, 5): ws.cell(r, j).border = BORD
                r += 1
            r += 1
            # топ-показатели по связи с целью
            ws.cell(r, 1, "Показатели, сильнее всего связанные с результатом:").font = Font(italic=True); r += 1
            _hrow(ws, r, ["показатель", "r с целью", "эффект", "категория",
                          "что это (для врача)", "p", "n"]); r += 1
            for _, row in ranked.head(10).iterrows():
                cat, desc = describe(row["показатель"])
                ws.cell(r, 1, _short(row["показатель"], 40))
                ws.cell(r, 2, round(float(row["r"]), 3))
                ws.cell(r, 3, effect_for(row["r"], tname))
                ws.cell(r, 4, cat)
                ws.cell(r, 5, desc).alignment = Alignment(wrap_text=True)
                pc = ws.cell(r, 6, round(float(row["p"]), 4))
                ws.cell(r, 7, int(row["n"]))
                if row["p"] < 0.01: pc.fill = SIG
                elif row["p"] < 0.05: pc.fill = SIG5
                for j in range(1, 8): ws.cell(r, j).border = BORD
                r += 1
            r += 2
        ws.column_dimensions["A"].width = 34
        ws.column_dimensions["D"].width = 24
        ws.column_dimensions["E"].width = 60
        for col in "BCFG": ws.column_dimensions[col].width = 12
    glossary_sheet(wb, df, used)
    write_methodology(wb, "regression")
    wb.save(out_path)
    return out_path


def pleiad_table_xlsx(out_path, df, targets, method="spearman", threshold=0.4,
                      top_k=12, alpha=None):
    """Табличная выгрузка связей плеяд (по каждой цели и кластеру)."""
    wb = Workbook(); wb.remove(wb.active)
    used = []
    for tname, target in targets.items():
        clusters = make_clusters(target)
        ws = wb.create_sheet(_short(tname, 28))
        r = 1
        for cl_name, mask in [("все", None), ("сильные", (clusters == "сильные").values),
                              ("слабые", (clusters == "слабые").values)]:
            _, edges_df, node_df = pleiad(df, target, cl_name, mask, method, threshold,
                                          top_k, alpha=alpha)
            plt.close("all")
            used += list(node_df["показатель"])
            ws.cell(r, 1, f"Кластер: {cl_name} — влияние на результат").font = Font(bold=True, size=11); r += 1
            _hrow(ws, r, ["показатель", "r с целью", "эффект", "категория",
                          "что это (для врача)", "p", "n"]); r += 1
            for _, row in node_df.iterrows():
                ws.cell(r, 1, _short(row["показатель"], 40))
                ws.cell(r, 2, round(float(row["r с целью"]), 3))
                ws.cell(r, 3, row["эффект"])
                ws.cell(r, 4, row["категория"])
                ws.cell(r, 5, row["что это (для врача)"]).alignment = Alignment(wrap_text=True)
                pc = ws.cell(r, 6, round(float(row["p"]), 4)); ws.cell(r, 7, int(row["n"]))
                if row["p"] < 0.01: pc.fill = SIG
                elif row["p"] < 0.05: pc.fill = SIG5
                for j in range(1, 8): ws.cell(r, j).border = BORD
                r += 1
            r += 2
        ws.column_dimensions["A"].width = 34
        ws.column_dimensions["D"].width = 24
        ws.column_dimensions["E"].width = 60
        for col in "BCFG": ws.column_dimensions[col].width = 12
    glossary_sheet(wb, df, used)
    write_methodology(wb, "pleiad")
    wb.save(out_path)
    return out_path


def glossary_sheet(wb, df, cols=None):
    """Лист «Словарь показателей». Если cols заданы — только эти столбцы
    (реально используемые в отчёте); иначе все."""
    if "Словарь показателей" in wb.sheetnames:
        return
    ws = wb.create_sheet("Словарь показателей")
    title = ("Словарь показателей — только используемые в этом отчёте "
             "(в уравнениях, ранжировании и плеядах)" if cols is not None
             else "Словарь показателей — простыми словами для врача/исследователя")
    ws.cell(1, 1, title).font = Font(bold=True, size=12)
    _hrow(ws, 2, ["Столбец", "Категория", "Что это и зачем"])
    r = 3
    res = set(point_columns(df))
    use = [c for c in dict.fromkeys(cols) if c in df.columns] if cols is not None else list(df.columns)
    for c in use:
        cat, desc = describe(c)
        if c in res:
            cat = "Результат"
        ws.cell(r, 1, str(c).replace("\n", " "))
        ws.cell(r, 2, cat)
        ws.cell(r, 3, desc).alignment = Alignment(wrap_text=True)
        for j in range(1, 4): ws.cell(r, j).border = BORD
        r += 1
    ws.freeze_panes = "A3"
    ws.column_dimensions["A"].width = 34
    ws.column_dimensions["B"].width = 30
    ws.column_dimensions["C"].width = 80
    return ws
