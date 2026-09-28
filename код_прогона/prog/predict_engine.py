# -*- coding: utf-8 -*-
"""Сохранение обученных моделей и применение к новым спортсменам.
Каждая модель сохраняется как «бандл» (.joblib) с паспортом: ID, что предсказывает,
какие признаки нужны на входе, ориентация результата, качество (train/CV R²).
Линейные модели дополнительно имеют готовую формулу — их можно считать вручную."""
import re
import numpy as np
import pandas as pd
import joblib
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment

import spearman_core as sc
from spearman_core import _to_float
from glossary import describe
from pleiad_engine import (point_columns, build_targets, make_clusters, _short,
                           equation_str, _ols, glossary_sheet)
from report_text import write_methodology
import ml_engine as ml

SCOPES = {"все игроки": None, "кластер: сильные": "сильные", "кластер: слабые": "слабые"}


def _target_series(df, target_name):
    tg, _ = build_targets(df)
    if target_name in tg:
        return tg[target_name]
    if target_name in df.columns:
        return df[target_name]
    raise ValueError(f"Цель «{target_name}» не найдена.")


def build_bundle(df, target_name, scope="все игроки", model_name="ГрадБустинг",
                 k=8, method="spearman", pooled=False):
    """Обучает выбранную модель и собирает бандл с паспортом."""
    orient = sc.orientation()
    if pooled:
        fit = ml.pooled_fit_one(df, model_name, k, method)
        mid = ml.make_model_id("pooled", "pooled", model_name)
        bundle = {
            "kind": "pooled", "model_id": mid, "model_name": model_name,
            "target": "попадания с точки (pooled)", "scope": "pooled",
            "orientation": orient, "features": fit["features"], "Xcols": fit["Xcols"],
            "points": fit["pts"], "model": fit["model"],
            "train_r2": fit["train_r2"], "cv_r2": fit["cv_r2"], "n": fit["n"],
        }
    else:
        tser = _target_series(df, target_name)
        mask = None
        if SCOPES.get(scope):
            mask = (make_clusters(tser) == SCOPES[scope]).values
        fit = ml.fit_one_model(df, tser, mask, model_name, k, method)
        mid = ml.make_model_id(target_name, scope, model_name)
        bundle = {
            "kind": "single", "model_id": mid, "model_name": model_name,
            "target": target_name, "scope": scope, "orientation": orient,
            "features": fit["features"], "model": fit["model"],
            "train_r2": fit["train_r2"], "cv_r2": fit["cv_r2"], "n": fit["n"],
        }
        # готовая формула для линейной
        if model_name == "Линейная":
            idx = df.index if mask is None else df.index[mask]
            data = pd.concat([tser.loc[idx].rename("__y__"),
                              df.loc[idx, fit["features"]]], axis=1).dropna()
            m = _ols(data["__y__"].values.astype(float),
                     data[fit["features"]].values.astype(float))
            bundle["formula"] = equation_str(target_name, fit["features"], m)
            bundle["coefs"] = dict(zip(["(константа)"] + fit["features"],
                                       [float(x) for x in m["beta"]]))
    return bundle


def save_bundle(bundle, out_path):
    joblib.dump(bundle, out_path)
    return out_path


def load_bundle(path):
    return joblib.load(path)


def input_template(bundle):
    """DataFrame-шаблон нужных входных столбцов (одна пустая строка для заполнения)."""
    cols = list(bundle["features"])
    if bundle["kind"] == "pooled":
        cols = cols + ["кластер (сильный/слабый)", "номер точки (1..N)"]
    tmpl = pd.DataFrame([{c: None for c in cols}])
    return tmpl


def passport_xlsx(bundle, out_path, template=True):
    """Excel-«паспорт» модели: ID, что предсказывает, входные признаки с описанием,
    качество, формула (если есть) + лист-шаблон для ввода новых спортсменов."""
    wb = Workbook(); ws = wb.active; ws.title = "Паспорт модели"
    bold = Font(bold=True)
    r = 1
    def row(label, value):
        nonlocal r
        ws.cell(r, 1, label).font = bold
        ws.cell(r, 2, value).alignment = Alignment(wrap_text=True)
        r += 1
    ws.cell(r, 1, "ПАСПОРТ МОДЕЛИ").font = Font(bold=True, size=14); r += 2
    row("ID модели", bundle["model_id"])
    row("Что предсказывает", bundle["target"])
    row("Тип модели", bundle["model_name"])
    row("Охват (на ком обучена)", bundle["scope"])
    row("Ориентация результата", "больше = лучше" if bundle["orientation"] == "high"
        else "меньше = лучше (время)")
    row("Обучено на N наблюдений", bundle["n"])
    row("Качество: R² на обучении", bundle["train_r2"])
    row("Качество: R² на кросс-валидации (честное)", bundle["cv_r2"])
    if bundle["cv_r2"] < 0.2:
        row("⚠ Предупреждение", "Низкий CV R² — прогноз ненадёжен, использовать с осторожностью.")
    if bundle.get("formula"):
        row("Готовая формула (можно считать вручную)", bundle["formula"])
    r += 1
    ws.cell(r, 1, "ВХОДНЫЕ ПРИЗНАКИ (что подать на вход):").font = Font(bold=True, size=12); r += 1
    ws.cell(r, 1, "Признак").font = bold
    ws.cell(r, 2, "Что это (для врача)").font = bold
    ws.cell(r, 3, "Категория").font = bold; r += 1
    for f in bundle["features"]:
        cat, desc = describe(f)
        ws.cell(r, 1, str(f).replace("\n", " "))
        ws.cell(r, 2, desc).alignment = Alignment(wrap_text=True)
        ws.cell(r, 3, cat); r += 1
    if bundle["kind"] == "pooled":
        r += 1
        ws.cell(r, 1, "Дополнительно для pooled: укажите кластер (сильный/слабый) и "
                      "номер точки (1..N).").font = Font(italic=True); r += 1
    ws.column_dimensions["A"].width = 40
    ws.column_dimensions["B"].width = 70
    ws.column_dimensions["C"].width = 26

    if template:
        wt = wb.create_sheet("Шаблон ввода")
        tmpl = input_template(bundle)
        for j, c in enumerate(tmpl.columns, start=1):
            wt.cell(1, j, str(c).replace("\n", " ")).font = bold
            wt.column_dimensions[wt.cell(1, j).column_letter].width = 22
        wt.cell(2, 1, "").value = None
        wt.cell(4, 1, "↑ Заполните значения строками (по одному спортсмену в строке), "
                      "сохраните и загрузите на шаге предсказания.").font = Font(italic=True)
    wb.save(out_path)
    return out_path


def _prep_features(df_new, feats):
    """Приводит входные столбцы к числам (запятые→точки), в нужном порядке."""
    X = pd.DataFrame()
    for f in feats:
        if f not in df_new.columns:
            raise ValueError(f"В данных нет столбца-признака: «{f}». "
                             f"Проверьте шаблон/названия.")
        X[f] = df_new[f].map(_to_float)
    return X


def predict(bundle, df_new):
    """Предсказание для новых спортсменов. df_new — таблица со столбцами-признаками
    (для pooled — плюс 'кластер (сильный/слабый)' и 'номер точки (1..N)').
    Возвращает копию df_new со столбцом 'Прогноз'."""
    feats = bundle["features"]
    out = df_new.copy()
    if bundle["kind"] == "single":
        X = _prep_features(df_new, feats)
        pred = bundle["model"].predict(X.values.astype(float))
    else:
        X = _prep_features(df_new, feats)
        # кластер
        cl_col = [c for c in df_new.columns if "кластер" in str(c).lower()]
        pt_col = [c for c in df_new.columns if "точк" in str(c).lower()]
        if not cl_col or not pt_col:
            raise ValueError("Для pooled нужны столбцы 'кластер (сильный/слабый)' и "
                             "'номер точки (1..N)'.")
        clv = df_new[cl_col[0]].astype(str).str.lower().str.startswith("с").astype(int)
        ptv = df_new[pt_col[0]].map(_to_float).fillna(1).astype(int)
        rows = []
        for i in range(len(df_new)):
            row = {c: 0.0 for c in bundle["Xcols"]}
            for f in feats:
                row[f] = X.iloc[i][f]
            row["cluster"] = float(clv.iloc[i])
            dum = f"точка_{int(ptv.iloc[i]) - 1}"      # one-hot как при обучении (drop_first)
            if dum in row:
                row[dum] = 1.0
            rows.append([row[c] for c in bundle["Xcols"]])
        pred = bundle["model"].predict(np.array(rows, dtype=float))
    out["Прогноз"] = np.round(pred, 2)
    return out
