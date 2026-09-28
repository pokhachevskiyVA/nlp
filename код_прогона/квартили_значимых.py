# -*- coding: utf-8 -*-
"""К каждому отчёту — отдельный файл с квартилями ТОЛЬКО значимых показателей.

Зачем: эксперт выуживал их из общей описательной статистики полтора часа.
Считается по `_вход.xlsx` той же подгруппы — это ровно те данные, на которых
считалась корреляция, поэтому числа гарантированно соответствуют отчёту.
Из общей описательной статистики брать нельзя для кластеров: «сильные» и «слабые»
в прогонах делятся по медиане СВОЕЙ цели, поэтому состав кластеров у разных целей разный.
"""
import sys, os, glob
sys.path.insert(0,'/Users/mac-os/Documents/Батя/outputs/_cache')
sys.path.insert(0,'/Users/mac-os/Documents/Батя/outputs/_cache/prog')
import numpy as np, pandas as pd
from glossary import describe
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

def сделать(папка, тихо=False):
    зп=os.path.join(папка,'Значимые_признаки.xlsx')
    вх=os.path.join(папка,'_вход.xlsx')
    if not (os.path.exists(зп) and os.path.exists(вх)): return None
    d=pd.read_excel(вх)
    имя_цели=d.columns[1]
    Z=pd.ExcelFile(зп)
    знач=Z.parse('Все значимые')
    без=Z.parse('Без повторов маркеров') if 'Без повторов маркеров' in Z.sheet_names else None
    пл1=Z.parse('Плеяда 1') if 'Плеяда 1' in Z.sheet_names else None
    пл2=Z.parse('Плеяда 2') if 'Плеяда 2' in Z.sheet_names else None
    в_пл1=set(пл1['признак'].astype(str)) if пл1 is not None else set()
    в_пл2=set(пл2['признак'].astype(str)) if пл2 is not None else set()
    в_без=set(без['признак'].astype(str)) if без is not None else set()

    def кв(s):
        s=pd.to_numeric(s,errors='coerce').dropna()
        if len(s)==0: return None
        return dict(n=len(s),**{'25Пц':s.quantile(.25),'Ме':s.median(),'75Пц':s.quantile(.75)})
    def зн(м):
        м=abs(м) if np.isfinite(м) else 1
        return 0 if м>=1000 else (1 if м>=100 else (2 if м>=1 else 3))

    стр=[]
    ц=кв(d[имя_цели])
    if ц:
        z=зн(ц['Ме'])
        стр.append(dict(признак='ЦЕЛЬ: '+str(имя_цели),система='— цель —',
                        n=ц['n'],**{k:round(ц[k],3) for k in ('25Пц','Ме','75Пц')},
                        **{'Ме [25Пц; 75Пц]':'%s [%s; %s]'%tuple(('%.*f'%(z,ц[k])).replace('.',',') for k in ('Ме','25Пц','75Пц'))},
                        r=np.nan,p=np.nan,плеяда='',описание='Целевая переменная этого прогона'))
    for _,r in знач.iterrows():
        c=str(r['признак'])
        if c not in d.columns: continue
        s=кв(d[c])
        if not s: continue
        кат,оп=describe(c)
        z=зн(s['Ме'])
        метки=[]
        if c in в_пл1: метки.append('плеяда 1')
        if c in в_пл2: метки.append('плеяда 2')
        if c in в_без: метки.append('без повторов')
        стр.append(dict(признак=c,система=кат,n=s['n'],
                        **{k:round(s[k],3) for k in ('25Пц','Ме','75Пц')},
                        **{'Ме [25Пц; 75Пц]':'%s [%s; %s]'%tuple(('%.*f'%(z,s[k])).replace('.',',') for k in ('Ме','25Пц','75Пц'))},
                        r=round(float(r['r']),3),p=float(r['p']),
                        плеяда=', '.join(метки),описание=оп))
    T=pd.DataFrame(стр)[['признак','система','n','25Пц','Ме','75Пц','Ме [25Пц; 75Пц]','r','p','плеяда','описание']]
    файл=os.path.join(папка,'Квартили_значимых.xlsx')
    with pd.ExcelWriter(файл,engine='openpyxl') as w:
        T.to_excel(w,sheet_name='Квартили значимых',index=False)
    # оформление
    wb=openpyxl.load_workbook(файл); ws=wb.active
    шапка=PatternFill('solid',fgColor='14315E'); бел=Font(color='FFFFFF',bold=True)
    цель_ф=PatternFill('solid',fgColor='FFF2CC'); пл_ф=PatternFill('solid',fgColor='E2EFDA')
    тонк=Side(style='thin',color='D5DBE3')
    for c in ws[1]: c.fill=шапка; c.font=бел; c.alignment=Alignment(horizontal='center',wrap_text=True)
    for row in ws.iter_rows(min_row=2):
        плеяда=str(row[9].value or '')
        for c in row:
            c.border=Border(left=тонк,right=тонк,top=тонк,bottom=тонк)
            if isinstance(c.value,float): c.number_format='0.000'
        if str(row[1].value)=='— цель —':
            for c in row: c.fill=цель_ф; c.font=Font(bold=True)
        elif 'плеяда 1' in плеяда:
            for c in row: c.fill=пл_ф
    ширины={'A':34,'B':26,'C':6,'D':13,'E':13,'F':13,'G':26,'H':8,'I':11,'J':22,'K':70}
    for к,ш in ширины.items(): ws.column_dimensions[к].width=ш
    ws.freeze_panes='B2'; ws.auto_filter.ref=ws.dimensions
    wb.save(файл)
    if not тихо: print("   %-58s значимых %2d"%(папка.split('ФИНАЛ_ВСЕ/')[-1],len(T)-1))
    return len(T)-1

if __name__=='__main__':
    import cpet
    корни=[f'{cpet.Y}/ГазоваяЭргометрия_2026-08-24',
           f'{cpet.Y}/ГазоваяЭргометрия_Контроль_2026-08-25']
    всего=0
    for к in корни:
        print("===",os.path.basename(к))
        for p in sorted(glob.glob(f'{к}/*/*/*')):
            if os.path.isdir(p):
                n=сделать(p)
                if n is not None: всего+=1
    print("\nфайлов создано:",всего)
