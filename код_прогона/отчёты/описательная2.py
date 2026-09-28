# -*- coding: utf-8 -*-
"""Описательная статистика по ВСЕМ признакам текущей таблицы, включая семейства
газ_коноо_*, газ_восст_ЧСС<125_* и КИ_*, которых не было в версии от 24 августа
(она строилась на таблице 2026-08-14, а эти признаки появились позже)."""
import sys, os, io
sys.path.insert(0,'/Users/mac-os/Documents/Батя/outputs/_cache')
sys.path.insert(0,'/Users/mac-os/Documents/Батя/outputs/_cache/prog')
import numpy as np, pandas as pd
from scipy import stats as st
from glossary import describe
import cpet
B='/Users/mac-os/Documents/Батя'
OUT=f'{cpet.Y}/Описательная_статистика_2026-08-27'
os.makedirs(OUT,exist_ok=True)
ПРИР=f'{cpet.Y}/ГазоваяЭргометрия_2026-08-24/Прирост'
xs=pd.ExcelFile(cpet.FULL)

# --- какие признаки участвуют в корреляции (тот же фильтр, что в прогонах) ---
пр_лыж=cpet.pick_features(xs.parse('Лыжники'),'газ','прирост')
пр_баск=cpet.pick_features(xs.parse('Баскетболисты'),'газ','прирост')
пр_к=cpet.pick_features(xs.parse('Контроль'),'газ','прирост')
нужны=set(пр_лыж)|set(пр_баск)|set(пр_к)
# порядок: сначала как в листе лыжников, потом то, что есть только у других
ПРИЗН=[c for c in xs.parse('Лыжники').columns if c in нужны]
for л in ('Баскетболисты','Контроль'):
    ПРИЗН+= [c for c in xs.parse(л).columns if c in нужны and c not in ПРИЗН]
# плюс служебные величины, которые эксперт тоже просил в цифрах
ДОП=[c for c in ('мин_нагрузки','КИ_за_нагрузку','КИ_за_разгон','Ннак_N1','Ннак_N2','Ннак_N3')
     if c in xs.parse('Лыжники').columns]
ВСЕ_К=ПРИЗН+[c for c in ДОП if c not in ПРИЗН]
print("признаков в корреляции:",len(ПРИЗН),"| плюс служебных:",len(ВСЕ_К)-len(ПРИЗН))
for сем in ('коноо','ЧСС<125','КИ_'):
    print("   семейство %-9s %d признаков"%(сем,len([c for c in ВСЕ_К if сем in str(c)])))

# --- группы ---
КЛ={}
for с in ('Лыжники','Баскетболисты'):
    v=pd.read_excel(f'{ПРИР}/{с}/Все/_вход.xlsx')
    КЛ[с]=dict(zip(v.iloc[:,0].astype(str),v['кластер']))
ГРУППЫ=[]
for с in ('Лыжники','Баскетболисты'):
    d=xs.parse(с); ф=d.iloc[:,0].astype(str)
    ГРУППЫ.append((f'{с}: все',d))
    for к in ('сильные','слабые'):
        ГРУППЫ.append((f'{с}: {к}',d[[КЛ[с].get(i)==к for i in ф]].reset_index(drop=True)))
ГРУППЫ.append(('Контроль',xs.parse('Контроль')))

def производные(d):
    """Величины, которых нет отдельной колонкой: прирост, кислородная и хронотропная цена."""
    T=pd.to_numeric(d['мин_нагрузки'],errors='coerce')
    Δ=30*(T.where(T>3)-3)
    t4=pd.to_numeric(d['газ_т4_t'],errors='coerce'); tk=pd.to_numeric(d['газ_коноо_t'],errors='coerce')
    V4=pd.to_numeric(d['газ_т4_VO2'],errors='coerce'); Vk=pd.to_numeric(d['газ_коноо_VO2'],errors='coerce')
    dW=30*(t4-tk)/60
    N=pd.to_numeric(d['КИ_за_разгон'],errors='coerce')
    Nн=pd.to_numeric(d['КИ_за_нагрузку'],errors='coerce')
    out=pd.DataFrame({
      'ЦЕЛЬ: Прирост нагрузки, Вт':Δ,
      'ЦЕЛЬ: Кислородная цена прироста, мл/мин/Вт':((V4-Vk)/dW.where(dW>15)).where(lambda v:v.between(0,40)),
      'ЦЕЛЬ: Хронотропная цена прироста, уд/Вт':(N/Δ.where(Δ>30)).where(lambda v:v.between(0,30)),
      'Средняя ЧСС за нагрузку, уд/мин':Nн/T,
      'Прирост мощности между коноо и т4, Вт':dW.where(dW>15),
    })
    пик=[c for c in d.columns if 'ик нагрузки' in str(c)]
    if пик: out['Пик нагрузки из таблиц заказчика, Вт']=pd.to_numeric(d[пик[0]],errors='coerce')
    return out


def стат(s):
    s=pd.to_numeric(s,errors='coerce').dropna()
    if len(s)==0: return dict(n=0)
    return dict(n=len(s),**{'25Пц':s.quantile(.25),'Ме':s.median(),'75Пц':s.quantile(.75),
                'IQR':s.quantile(.75)-s.quantile(.25),'мин':s.min(),'макс':s.max(),
                'среднее':s.mean(),'СО':s.std(ddof=1) if len(s)>1 else np.nan})
def ф(v,z):
    if not np.isfinite(v): return "—"
    return ("%.*f"%(z,v)).replace(".",",")
def зн(м):
    м=abs(м) if np.isfinite(м) else 1
    return 0 if м>=1000 else (1 if м>=100 else (2 if м>=1 else 3))

# --- листы детализации ---
листы={}
for имя,d in ГРУППЫ:
    стр=[]
    пров=производные(d)
    for c in пров.columns:
        s=стат(пров[c])
        if s['n']: стр.append(dict(признак=c,система='Производная величина',
                                   описание='Считается из других столбцов, отдельной колонки в таблице нет',**s))
    for c in ВСЕ_К:
        if c not in d.columns: continue
        s=стат(d[c])
        if s['n']==0: continue
        кат,оп=describe(c)
        стр.append(dict(признак=c,система=кат,описание=оп,**s))
    t=pd.DataFrame(стр)
    листы[имя.replace(': ','_')]=t
# --- сводные в формате Ме [25;75] + Краскел-Уоллис ---
def свод(набор,подпись):
    стр=[]
    прод={и:производные(d) for и,d in набор}
    список=list(прод[набор[0][0]].columns)+list(ВСЕ_К)
    for c in список:
        ряды={}
        for имя,d in набор:
            src=прод[имя] if c in прод[имя].columns else d
            if c in src.columns: ряды[имя]=pd.to_numeric(src[c],errors='coerce').dropna()
        if not ряды or max(len(v) for v in ряды.values())==0: continue
        м0=list(ряды.values())[0].median() if len(list(ряды.values())[0]) else np.nan
        z=зн(м0)
        кат,оп=describe(c)
        р=dict(признак=c,система=кат)
        for имя in [и for и,_ in набор]:
            v=ряды.get(имя,pd.Series(dtype=float))
            р[имя+' n']=len(v)
            р[имя]="%s [%s; %s]"%(ф(v.median(),z),ф(v.quantile(.25),z),ф(v.quantile(.75),z)) if len(v) else "—"
        сп=[v for v in ряды.values() if len(v)>3]
        р['p, Краскел–Уоллис']=st.kruskal(*сп).pvalue if len(сп)>=3 else np.nan
        р['описание']=оп
        стр.append(р)
    t=pd.DataFrame(стр)
    кол=['признак','система']+[c for и,_ in набор for c in (и,и+' n')]+['p, Краскел–Уоллис','описание']
    return t[кол]

ГР=dict(ГРУППЫ)
своды={
 'СВОД_три_группы':свод([('Лыжники: все',ГР['Лыжники: все']),('Баскетболисты: все',ГР['Баскетболисты: все']),
                         ('Контроль',ГР['Контроль'])],'три группы'),
 'СВОД_Лыжники_кластеры':свод([('Лыжники: сильные',ГР['Лыжники: сильные']),
                               ('Лыжники: слабые',ГР['Лыжники: слабые']),('Контроль',ГР['Контроль'])],'лыжники'),
 'СВОД_Баскет_кластеры':свод([('Баскетболисты: сильные',ГР['Баскетболисты: сильные']),
                              ('Баскетболисты: слабые',ГР['Баскетболисты: слабые']),('Контроль',ГР['Контроль'])],'баскет'),
}
длинная=[]
for имя,t in листы.items():
    tt=t.copy(); tt.insert(0,'группа',имя); длинная.append(tt)
длинная=pd.concat(длинная,ignore_index=True)

файл=f'{OUT}/Описательная_статистика_2026-08-27.xlsx'
with pd.ExcelWriter(файл,engine='openpyxl') as w:
    for имя,t in своды.items(): t.to_excel(w,sheet_name=имя[:31],index=False)
    for имя,t in листы.items(): t.to_excel(w,sheet_name=имя[:31],index=False)
    длинная.to_excel(w,sheet_name='ВСЁ_одной_таблицей',index=False)
import openpyxl
wb=openpyxl.load_workbook(файл)
for ws in wb.worksheets:
    for col in ws.columns:
        L=max((len(str(c.value)) for c in col if c.value is not None),default=8)
        ws.column_dimensions[col[0].column_letter].width=min(max(L+2,10),48)
    ws.freeze_panes='C2'
    for r in ws.iter_rows(min_row=2):
        for c in r:
            if isinstance(c.value,float): c.number_format='0.000'
wb.save(файл)
print("\nзаписано:",файл)
print("листов:",len(wb.sheetnames))
print("строк в длинной:",len(длинная))
k=своды['СВОД_три_группы']
print("\nпроверка — семейство коноо в своде:",len([c for c in k['признак'] if 'коноо' in str(c)]),"признаков")
print(k[k['признак'].astype(str).str.contains('коноо')][['признак','Лыжники: все','Баскетболисты: все','Контроль','p, Краскел–Уоллис']].to_string(index=False))
